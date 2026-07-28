#!/usr/bin/env python3
"""Tests for report_export.py — the shared report_data -> row projection.

The load-bearing guarantee is COLUMN PARITY: xlsx, CSV and XML must project the
same findings the same way, because report_ingest.py is written as the inverse of
that one projection. A test that only checked CSV would let the xlsx drift back
apart, which is the exact duplication this module was created to remove.
"""
import io
import json
import os
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

import report_export as RE  # noqa: E402
import report_ingest as RI  # noqa: E402

try:
    import report_xlsx_build as RX
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAVE_OPENPYXL = False

DATA = {
    "engagement": {"subtitle": "Example — API", "target": "api.example.com",
                   "date": "2026-01-01", "classification": "CONFIDENTIAL"},
    "findings": [
        {"id": "F-002", "title": "Outdated component", "severity": "Medium",
         "cvss_score": 5.3, "cvss_vector": "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:L/I:N/A:N",
         "cwe": "CWE-1104", "cves": [{"id": "CVE-2021-23017"}],
         "attack": ["AML.T0051"], "affected": ["api.example.com"]},
        {"id": "F-001", "title": "Broken object level authorization", "severity": "High",
         "cvss_score": 7.5, "cwe": "CWE-639", "owasp": "API1:2023",
         "attack": ["T1190", "T1059.001"], "affected": ["GET /orders/{id}"],
         "references": ["https://owasp.org/API-Security/"]},
        {"id": "F-003", "title": "Info leak", "severity": "Info",
         "needs_live_confirmation": True},
    ],
}


def _csv_rows(text):
    import csv
    return list(csv.reader(io.StringIO(text.lstrip("﻿"), newline="")))


class TestNeutralize(unittest.TestCase):
    def test_every_formula_lead_is_disarmed(self):
        for lead in ("=", "+", "-", "@", "\t", "\r"):
            payload = lead + "cmd|'/c calc'!A1"
            self.assertTrue(RE.neutralize(payload).startswith("'"), lead)

    def test_ordinary_values_are_untouched(self):
        for v in ("F-001", "Broken auth", "CVSS:3.1/AV:N", "", "3.5", "—"):
            self.assertEqual(RE.neutralize(v), v)

    def test_read_cell_is_the_exact_inverse(self):
        for v in ("=1+1", "-lead", "@at", "+plus", "normal", "", "it's fine", "'quoted"):
            self.assertEqual(RE.read_cell(RE.neutralize(v)), v)

    def test_read_cell_leaves_a_genuine_leading_quote_alone(self):
        # "'hello" is a real value, not a neutralized one: char 2 is not a
        # formula lead, so stripping the quote would corrupt it.
        self.assertEqual(RE.read_cell("'hello"), "'hello")


class TestProjection(unittest.TestCase):
    def test_attack_accepts_list_string_and_atlas(self):
        self.assertEqual(RE.attack_str({"attack": ["T1190", "T1059.001"]}), "T1190, T1059.001")
        self.assertEqual(RE.attack_str({"attack": "AML.T0051"}), "AML.T0051")
        self.assertEqual(RE.attack_str({}), "")

    def test_rows_are_severity_major_then_id(self):
        ids = [RE.finding_row(f)[0] for f in sorted(RE.findings_of(DATA), key=RE.sort_key)]
        self.assertEqual(ids, ["F-001", "F-002", "F-003"])   # High, Medium, Info

    def test_row_width_matches_headers(self):
        for f in RE.findings_of(DATA):
            self.assertEqual(len(RE.finding_row(f)), len(RE.HEADERS))


class TestCsv(unittest.TestCase):
    def setUp(self):
        self.text = RE.to_csv(DATA)

    def test_has_utf8_bom_and_crlf(self):
        # Without the BOM, Excel reads UTF-8 as the local codepage and mangles
        # the em-dash placeholders this projection emits.
        self.assertTrue(self.text.startswith("﻿"))
        self.assertIn("\r\n", self.text)

    def test_header_row_is_the_shared_headers(self):
        self.assertEqual(_csv_rows(self.text)[0], RE.HEADERS)

    def test_one_row_per_finding_in_severity_order(self):
        rows = _csv_rows(self.text)[1:]
        self.assertEqual(len(rows), 3)
        self.assertEqual([r[0] for r in rows], ["F-001", "F-002", "F-003"])

    def test_mitre_column_carries_attack_and_atlas(self):
        i = RE.HEADERS.index("MITRE")
        rows = {r[0]: r for r in _csv_rows(self.text)[1:]}
        self.assertEqual(rows["F-001"][i], "T1190, T1059.001")
        self.assertEqual(rows["F-002"][i], "AML.T0051")
        self.assertEqual(rows["F-003"][i], "—")

    def test_embedded_comma_and_newline_survive_quoting(self):
        d = {"findings": [{"id": "X", "title": 'a,b "q"\nsecond', "severity": "Low"}]}
        row = _csv_rows(RE.to_csv(d))[1]
        self.assertEqual(row[1], 'a,b "q"\nsecond')


class TestXml(unittest.TestCase):
    def test_is_well_formed_and_mirrors_the_rows(self):
        root = ET.fromstring(RE.to_xml(DATA))
        fs = root.find("findings")
        self.assertEqual(fs.get("count"), "3")
        self.assertEqual([f.findtext("id") for f in fs], ["F-001", "F-002", "F-003"])

    def test_element_names_match_the_shared_columns(self):
        want = [tag for _, tag, _ in RE.REGISTER_COLUMNS]
        root = ET.fromstring(RE.to_xml(DATA))
        got = [c.tag for c in root.find("findings")[0]]
        self.assertEqual(got, want)

    def test_markup_in_a_value_is_escaped_not_injected(self):
        d = {"findings": [{"id": "X", "title": "<script>alert(1)</script>&", "severity": "Low"}]}
        raw = RE.to_xml(d)
        self.assertNotIn("<script>", raw)
        self.assertEqual(ET.fromstring(raw).find("findings/finding/title").text,
                         "<script>alert(1)</script>&")


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
class TestColumnParity(unittest.TestCase):
    """The reason this module exists: one projection, three renderings."""

    def test_xlsx_register_header_row_is_the_shared_headers(self):
        wb, _ = RX.build_workbook(DATA, {"target": None, "title": None,
                                         "location": None, "date": None})
        ws = wb["Findings Register"]
        got = [ws.cell(row=7, column=i + 1).value for i in range(len(RE.HEADERS))]
        self.assertEqual(got, RE.HEADERS)

    def test_xlsx_and_csv_agree_cell_for_cell(self):
        wb, _ = RX.build_workbook(DATA, {"target": None, "title": None,
                                         "location": None, "date": None})
        ws = wb["Findings Register"]
        csv_rows = _csv_rows(RE.to_csv(DATA))[1:]
        for r, want in enumerate(csv_rows, start=8):
            got = [RE.read_cell(ws.cell(row=r, column=c + 1).value)
                   for c in range(len(RE.HEADERS))]
            self.assertEqual(got, [RE.read_cell(v) for v in want], f"row {r}")

    def test_a_formula_title_round_trips_through_xlsx_and_ingest(self):
        d = json.loads(json.dumps(DATA))
        d["findings"][1]["title"] = "=cmd|'/c calc'!A1"
        wb, _ = RX.build_workbook(d, {"target": None, "title": None,
                                      "location": None, "date": None})
        with tempfile.TemporaryDirectory() as t:
            p = os.path.join(t, "r.xlsx")
            wb.save(p)
            res, code = RI.ingest(p)
        self.assertEqual(code, RI.EXIT_OK)
        got = {f["id"]: f for f in res["findings"]}
        self.assertEqual(got["F-001"]["title"], "=cmd|'/c calc'!A1")
        self.assertEqual(got["F-001"]["attack"], ["T1190", "T1059.001"])
        self.assertEqual(got["F-002"]["attack"], ["AML.T0051"])


class TestCli(unittest.TestCase):
    def test_writes_requested_format_and_reports_count(self):
        with tempfile.TemporaryDirectory() as t:
            src = os.path.join(t, "report_data.json")
            with open(src, "w", encoding="utf-8") as fh:
                json.dump(DATA, fh)
            for fmt, probe in (("csv", "ID,Finding"), ("xml", "<findings")):
                out = os.path.join(t, "o." + fmt)
                self.assertEqual(RE.main([src, "--format", fmt, "-o", out]), RE.EXIT_OK)
                self.assertIn(probe, open(out, encoding="utf-8").read())

    def test_missing_file_is_usage_error_not_traceback(self):
        import contextlib
        with contextlib.redirect_stderr(io.StringIO()):
            rc = RE.main([os.path.join(tempfile.gettempdir(), "nope.json"), "--format", "csv"])
        self.assertEqual(rc, RE.EXIT_USAGE)


if __name__ == "__main__":
    unittest.main(verbosity=2)
